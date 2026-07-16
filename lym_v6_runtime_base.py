from __future__ import annotations

"""
LYM AUTO CONTROL V6.0 - Actualización fiscal, gastos, fechas y reportería.

Este archivo debe permanecer en la misma carpeta que:
    LYM_AUTO_CONTROL_V5_0_LEASING.py
    lym_v6_core.py

Ejecutar:
    python LYM_AUTO_CONTROL_V6_0_FINAL.py

La actualización se monta de forma no destructiva sobre V5 para conservar
compatibilidad con los registros cifrados y carpetas ya existentes.
"""

import re
import sys
import tempfile
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Optional

import LYM_AUTO_CONTROL_V5_0_LEASING as base
from lym_v6_core import (
    DEFAULT_RETROACTIVE_DAYS,
    DUCA_SUBCATEGORY,
    DUCA_TAX_TYPES,
    PERM_FECHAS_RETROACTIVAS,
    calculate_duca_total,
    calculate_sale_price,
    expense_amounts,
    migrate_cost_item_v6,
    norm,
    normalize_duca_lines,
    summarize_vehicle_costs,
    validate_manual_price,
    validate_operational_date,
    workbook_validation_examples,
)

APP_VERSION_V6 = "6.0.0_FISCAL"
STANDARD_COLORS = [
    "BLANCO", "NEGRO", "GRIS", "PLATEADO", "ROJO", "AZUL", "VERDE",
    "AMARILLO", "NARANJA", "BEIGE", "CAFÉ", "MARRÓN", "DORADO",
    "MORADO", "CELESTE", "VINO", "CHAMPAGNE", "OTRO",
]

STANDARD_EXPENSES = [
    ("GRUA_TRASLADO_USA", "Grúa / traslado USA"),
    ("COSTO_EXTRA_TRANSITO", "Costo extra tránsito"),
    ("NAVIERA_GRUA_INTERNA", "Naviera - grúa interna"),
    ("NAVIERA_FLETE", "Naviera - flete"),
    ("NAVIERA_BL", "Naviera - BL"),
    ("IMPUESTOS_ADUANA", "Impuestos DUCA"),
    ("ALMACENAMIENTO_ADUANA", "Almacenamiento"),
    ("TRAMITE_ADUANAL", "Servicio trámite aduanal"),
    ("EMISIONES", "Costo emisiones"),
    ("CITA", "Costo cita"),
    ("PLACAS", "Costo placas"),
    ("ENDEREZADO_PINTURA", "Enderezado y pintura"),
    ("SERVICIOS_MECANICOS", "Servicios mecánicos"),
    ("GRUA_LOCAL", "Grúa local"),
    ("REPUESTO_DETALLE", "Repuestos"),
]


def _to_float(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(value)
    except Exception:
        return default


def _has_extended_date_permission(user: dict) -> bool:
    try:
        return bool(base.user_has_permission(user or {}, PERM_FECHAS_RETROACTIVAS))
    except Exception:
        return False


def _walk_dates(value: Any) -> set[str]:
    found: set[str] = set()
    if isinstance(value, dict):
        for item in value.values():
            found.update(_walk_dates(item))
    elif isinstance(value, (list, tuple)):
        for item in value:
            found.update(_walk_dates(item))
    elif isinstance(value, (date, datetime)):
        found.add(value.date().isoformat() if isinstance(value, datetime) else value.isoformat())
    elif isinstance(value, str):
        parsed = base._parse_date(value)
        if parsed:
            found.add(parsed.isoformat())
    return found


def _date_widget_is_active(widget) -> bool:
    parent = widget.parentWidget()
    while parent is not None:
        checkbox = getattr(parent, "chk", None)
        if checkbox is not None:
            return bool(checkbox.isChecked())
        parent = parent.parentWidget()
    return widget.isEnabled()


def _validate_dialog_dates(dialog, user: dict, existing_values: Iterable[Any] = ()) -> tuple[bool, str]:
    existing = set(existing_values)
    for edit in dialog.findChildren(base.QDateEdit):
        if not _date_widget_is_active(edit):
            continue
        selected = edit.date().toPython()
        ok, msg = validate_operational_date(
            selected,
            has_extended_permission=_has_extended_date_permission(user),
            existing_values=existing,
            max_back_days=DEFAULT_RETROACTIVE_DAYS,
        )
        if not ok:
            return False, msg
    return True, "OK"


def _report_output_dir() -> Path:
    try:
        return base._report_output_dir()
    except Exception:
        root = base.get_data_folder()
        if root is None:
            raise RuntimeError("No hay carpeta activa del sistema.")
        folder = Path(root) / base.SUB_REPORTES
        folder.mkdir(parents=True, exist_ok=True)
        return folder


def vehicle_total_cost_v6(vehicle: dict) -> float:
    return float(summarize_vehicle_costs(vehicle).get("net_vehicle_cost", 0.0))


def _store_new_cost_document(
    src: str,
    vehicle_code: str,
    doc_type: str,
    label: str,
    old_path: str = "",
    old_name: str = "",
) -> tuple[str, str]:
    if src and Path(src).exists():
        rel, name = base._store_cost_doc(src, vehicle_code, doc_type, label)
        if rel:
            return rel, name
    return old_path or "", old_name or ""


def sync_stage_cost_items_v6(
    vehicle: dict,
    stage_key: str,
    cost_items: list[dict],
    user: dict,
    device,
) -> float:
    """Guarda gastos con factura, OC, comprobante y crédito fiscal.

    El total devuelto es el costo NETO del vehículo. El desembolso bruto y el
    crédito fiscal quedan conservados por separado en cada registro.
    """
    base.ensure_vehicle_runtime_fields(vehicle)
    source = f"stage:{stage_key}"
    keep = [g for g in vehicle.get("gastos_detallados", []) if g.get("source") != source]
    total_net = 0.0

    for raw in cost_items or []:
        gross = round(max(_to_float(raw.get("monto_usd")), 0.0), 2)
        credit = round(max(_to_float(raw.get("credito_fiscal_usd")), 0.0), 2)
        duca_lines = normalize_duca_lines(raw.get("duca_tributos") or [])
        is_duca = (
            norm(raw.get("subcategoria")).replace(" ", "_") == DUCA_SUBCATEGORY
            or bool(duca_lines)
            or bool(raw.get("credito_fiscal_duca_usd"))
        )
        if is_duca:
            duca_total = calculate_duca_total(duca_lines)
            gross = round(max(gross, duca_total), 2)
            credit = round(max(credit, duca_total, _to_float(raw.get("credito_fiscal_duca_usd"))), 2)
            net = 0.0
        else:
            if credit > gross and gross > 0:
                credit = gross
            net = round(max(gross - credit, 0.0), 2)

        if gross <= 0 and credit <= 0:
            continue

        label = base._safe_filename(
            f"{raw.get('categoria', 'GASTO')}_{raw.get('subcategoria', 'DETALLE')}"
        )
        vehicle_code = vehicle.get("codigo", "VEHICULO")

        factura_rel, factura_name = _store_new_cost_document(
            raw.get("factura_src") or raw.get("duca_src") or "",
            vehicle_code,
            "FACTURA_DUCA" if is_duca else "FACTURA",
            label,
            raw.get("factura_documento", ""),
            raw.get("factura_documento_nombre", ""),
        )
        comp_rel, comp_name = _store_new_cost_document(
            raw.get("comprobante_src") or "",
            vehicle_code,
            "COMPROBANTE_PAGO",
            label,
            raw.get("comprobante", ""),
            raw.get("comprobante_nombre", ""),
        )
        oc_rel, oc_name = _store_new_cost_document(
            raw.get("oc_src") or "",
            vehicle_code,
            "OC",
            label,
            raw.get("oc_documento", ""),
            raw.get("oc_documento_nombre", ""),
        )

        item = {
            "id": raw.get("id") or uuid.uuid4().hex,
            "source": source,
            "stage_key": stage_key,
            "categoria": norm(raw.get("categoria") or stage_key),
            "subcategoria": norm(raw.get("subcategoria") or "GENERAL").replace(" ", "_"),
            "fecha": (base._parse_date(raw.get("fecha")) or date.today()).isoformat(),
            "descripcion": str(raw.get("descripcion") or "").strip(),
            "monto_usd": gross,
            "credito_fiscal_usd": credit,
            "credito_fiscal_duca_usd": credit if is_duca else 0.0,
            "costo_neto_usd": net,
            "tratamiento_fiscal": "DUCA_CREDITO_TOTAL" if is_duca else "FACTURA_CREDITO_SEPARADO",
            "duca_tributos": duca_lines,
            "duca_numero": str(raw.get("duca_numero") or "").strip().upper(),
            "proveedor": str(raw.get("proveedor") or "").strip(),
            "factura_numero": str(raw.get("factura_numero") or "").strip().upper(),
            "oc_numero": str(raw.get("oc_numero") or "").strip().upper(),
            "factura_documento": factura_rel,
            "factura_documento_nombre": factura_name,
            "oc_documento": oc_rel,
            "oc_documento_nombre": oc_name,
            "comprobante": comp_rel,
            "comprobante_nombre": comp_name,
            "comentario_gasto": str(raw.get("comentario_gasto") or "").strip(),
            "usuario": user.get("usuario", ""),
            "fecha_registro": base._now_iso(),
        }
        keep.append(item)
        total_net += net

    vehicle["gastos_detallados"] = keep
    return round(total_net, 2)


def _autosize_sheet(ws, min_width: int = 10, max_width: int = 34) -> None:
    from openpyxl.utils import get_column_letter

    for col in range(1, ws.max_column + 1):
        width = min_width
        for row in range(1, min(ws.max_row, 250) + 1):
            value = ws.cell(row, col).value
            width = max(width, len(str(value or "")) + 2)
        ws.column_dimensions[get_column_letter(col)].width = min(width, max_width)


def _style_report_sheet(ws, title: str, last_col: int, freeze: str = "A4") -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    navy = "08285A"
    orange = "F59A13"
    light = "EAF1FA"
    border = Border(
        left=Side(style="thin", color="D7E0EC"),
        right=Side(style="thin", color="D7E0EC"),
        top=Side(style="thin", color="D7E0EC"),
        bottom=Side(style="thin", color="D7E0EC"),
    )
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(size=17, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=orange)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[3].height = 42

    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=light)

    ws.freeze_panes = freeze
    if ws.max_row >= 3:
        ws.auto_filter.ref = f"A3:{get_column_letter(last_col)}{ws.max_row}"
