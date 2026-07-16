from __future__ import annotations

from lym_v6_runtime_base import *
from lym_v6_runtime_base import _autosize_sheet, _report_output_dir, _style_report_sheet, _to_float

def generate_spare_parts_excel_v6(vehicles: list[dict], user: dict):
    if not base.user_has_permission(user, base.PERM_GENERATE_REPORTS):
        return False, "No tienes permiso para generar reportes.", None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        out = _report_output_dir() / f"REPORTE_REPUESTOS_LYM_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Repuestos comprados"
        headers = [
            "Fecha", "Código vehículo", "Marca", "Modelo", "Año", "VIN / Chasis",
            "Descripción repuesto", "Proveedor", "Valor pagado USD",
            "Crédito fiscal USD", "Costo neto USD", "Número factura", "Número OC",
            "Factura PDF", "OC PDF", "Comprobante pago", "Usuario", "Comentario",
        ]
        ws.append([])
        ws.append([])
        ws.append(headers)

        count = 0
        for vehicle in vehicles:
            base.ensure_vehicle_runtime_fields(vehicle)
            for raw in vehicle.get("gastos_detallados", []) or []:
                if norm(raw.get("subcategoria")).replace(" ", "_") != "REPUESTO_DETALLE":
                    continue
                amounts = expense_amounts(raw)
                ws.append([
                    base._fmt_date(raw.get("fecha")),
                    vehicle.get("codigo", ""),
                    vehicle.get("marca", ""),
                    vehicle.get("modelo", ""),
                    vehicle.get("anio", ""),
                    vehicle.get("vin") or vehicle.get("chasis") or "",
                    raw.get("descripcion", ""),
                    raw.get("proveedor", ""),
                    amounts["gross"],
                    amounts["credit"],
                    amounts["net"],
                    raw.get("factura_numero", ""),
                    raw.get("oc_numero", ""),
                    raw.get("factura_documento_nombre") or raw.get("factura_documento") or "",
                    raw.get("oc_documento_nombre") or raw.get("oc_documento") or "",
                    raw.get("comprobante_nombre") or raw.get("comprobante") or "",
                    raw.get("usuario", ""),
                    raw.get("comentario_gasto", ""),
                ])
                count += 1

        data_end_row = ws.max_row
        total_row = data_end_row + 2
        ws.cell(total_row, 7, "TOTALES")
        ws.cell(total_row, 7).font = Font(bold=True)
        for col in (9, 10, 11):
            letter = ws.cell(3, col).column_letter
            if data_end_row >= 4:
                ws.cell(total_row, col, f"=SUM({letter}4:{letter}{data_end_row})")
            else:
                ws.cell(total_row, col, 0)
            ws.cell(total_row, col).number_format = '"$ "#,##0.00'

        for row in range(4, ws.max_row + 1):
            for col in (9, 10, 11):
                ws.cell(row, col).number_format = '"$ "#,##0.00'
        _style_report_sheet(ws, "L&M INVERSIONES - REPORTE DE REPUESTOS COMPRADOS", len(headers))
        _autosize_sheet(ws)
        wb.save(out)
        base.log_audit("REPORTE_REPUESTOS_V6", user.get("usuario", ""), "", f"{count} repuestos")
        return True, f"Reporte generado con {count} repuestos.", out
    except Exception as exc:
        return False, f"No se pudo generar el reporte de repuestos: {type(exc).__name__}: {exc}", None


def generate_horizontal_expense_excel_v6(vehicles: list[dict], user: dict):
    if not base.user_has_permission(user, base.PERM_GENERATE_REPORTS):
        return False, "No tienes permiso para generar reportes.", None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        out = _report_output_dir() / f"GASTOS_HORIZONTALES_CREDITO_FISCAL_LYM_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Gastos horizontales"

        fixed_headers = [
            "Código", "Marca", "Modelo", "Año", "VIN / Chasis", "Tipo compra",
            "Estado operativo", "Estado comercial", "Costo compra USD",
        ]
        expense_headers: list[str] = []
        for _, label in STANDARD_EXPENSES:
            expense_headers.extend([f"Gasto {label} USD", f"Crédito fiscal {label} USD"])
        duca_headers = [f"{tax} DUCA USD" for tax in DUCA_TAX_TYPES]
        final_headers = [
            "Total gastos pagados USD", "Crédito fiscal facturas USD",
            "Crédito fiscal DUCA USD", "Total crédito fiscal USD",
            "Costo neto vehículo USD", "Utilidad configurada %",
            "Subtotal venta sin IVA USD", "IVA venta 13% USD",
            "Precio final cliente USD",
        ]
        headers = fixed_headers + expense_headers + duca_headers + final_headers
        ws.append([])
        ws.append([])
        ws.append(headers)

        for vehicle in vehicles:
            base.ensure_vehicle_runtime_fields(vehicle)
            summary = summarize_vehicle_costs(vehicle)
            expense_values: list[float] = []
            for key, _ in STANDARD_EXPENSES:
                bucket = summary["expenses"].get(key, {})
                expense_values.extend([
                    _to_float(bucket.get("gross")),
                    _to_float(bucket.get("credit")),
                ])
            duca_values = [_to_float(summary["duca_taxes"].get(tax)) for tax in DUCA_TAX_TYPES]
            margin_pct = _to_float((vehicle.get("precio_final") or {}).get("margen_pct"), 15.0)
            price = calculate_sale_price(
                summary["gross_paid"],
                margin_pct,
                iva_pct=13,
                fiscal_credit=summary["total_fiscal_credit"],
            )
            row = [
                vehicle.get("codigo", ""),
                vehicle.get("marca", ""),
                vehicle.get("modelo", ""),
                vehicle.get("anio", ""),
                vehicle.get("vin") or vehicle.get("chasis") or "",
                vehicle.get("tipo_compra", ""),
                base.STAGE_META.get(vehicle.get("estado_actual"), {}).get("label", vehicle.get("estado_actual", "")),
                vehicle.get("estado_comercial", ""),
                summary["purchase_cost"],
            ] + expense_values + duca_values + [
                summary["gross_paid"],
                summary["generic_fiscal_credit"],
                summary["duca_fiscal_credit"],
                summary["total_fiscal_credit"],
                summary["net_vehicle_cost"],
                margin_pct,
                price["subtotal_before_iva"],
                price["iva_sale"],
                price["final_price"],
            ]
            ws.append(row)

        money_start = 9
        data_end_row = ws.max_row
        for row in range(4, data_end_row + 1):
            for col in range(money_start, len(headers) + 1):
                header = str(ws.cell(3, col).value or "")
                if header == "Utilidad configurada %":
                    ws.cell(row, col).number_format = '0.00'
                else:
                    ws.cell(row, col).number_format = '"$ "#,##0.00'

        total_row = data_end_row + 2
        ws.cell(total_row, 1, "TOTALES")
        ws.cell(total_row, 1).font = Font(bold=True)
        for col in range(money_start, len(headers) + 1):
            header = str(ws.cell(3, col).value or "")
            if header == "Utilidad configurada %":
                continue
            letter = ws.cell(3, col).column_letter
            if data_end_row >= 4:
                ws.cell(total_row, col, f"=SUM({letter}4:{letter}{data_end_row})")
            else:
                ws.cell(total_row, col, 0)
            ws.cell(total_row, col).number_format = '"$ "#,##0.00'

        _style_report_sheet(
            ws,
            "L&M INVERSIONES - GASTOS POR VEHÍCULO Y CRÉDITO FISCAL",
            len(headers),
        )
        _autosize_sheet(ws, max_width=28)

        detail = wb.create_sheet("Detalle de gastos")
        detail_headers = [
            "Código", "Vehículo", "Etapa", "Categoría", "Subcategoría", "Descripción",
            "Fecha", "Proveedor", "Factura", "OC", "Gasto pagado USD",
            "Crédito fiscal USD", "Crédito fiscal DUCA USD", "Costo neto USD",
            "Factura PDF", "OC PDF", "Comprobante pago", "Comentario",
        ]
        detail.append([])
        detail.append([])
        detail.append(detail_headers)
        for vehicle in vehicles:
            for raw in vehicle.get("gastos_detallados", []) or []:
                if not isinstance(raw, dict) or base._v50_is_purchase_cost_item(raw):
                    continue
                amounts = expense_amounts(raw)
                detail.append([
                    vehicle.get("codigo", ""),
                    f"{vehicle.get('marca', '')} {vehicle.get('modelo', '')} {vehicle.get('anio', '')}",
                    base.STAGE_META.get(raw.get("stage_key"), {}).get("label", raw.get("stage_key", "")),
                    raw.get("categoria", ""),
                    raw.get("subcategoria", ""),
                    raw.get("descripcion", ""),
                    base._fmt_date(raw.get("fecha")),
                    raw.get("proveedor", ""),
                    raw.get("factura_numero", ""),
                    raw.get("oc_numero", ""),
                    amounts["gross"],
                    amounts["credit"],
                    amounts["duca_credit"],
                    amounts["net"],
                    raw.get("factura_documento_nombre") or raw.get("factura_documento") or "",
                    raw.get("oc_documento_nombre") or raw.get("oc_documento") or "",
                    raw.get("comprobante_nombre") or raw.get("comprobante") or "",
                    raw.get("comentario_gasto", ""),
                ])
        for row in range(4, detail.max_row + 1):
            for col in (11, 12, 13, 14):
                detail.cell(row, col).number_format = '"$ "#,##0.00'
        _style_report_sheet(detail, "DETALLE DOCUMENTAL DE GASTOS", len(detail_headers))
        _autosize_sheet(detail, max_width=32)

        wb.save(out)
        base.log_audit("REPORTE_GASTOS_CREDITO_V6", user.get("usuario", ""), "", f"{len(vehicles)} vehículos")
        return True, f"Reporte horizontal generado con {len(vehicles)} vehículos.", out
    except Exception as exc:
        return False, f"No se pudo generar el reporte horizontal: {type(exc).__name__}: {exc}", None
